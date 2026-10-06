"""
Bundle verification tests for the PyInstaller-frozen StrictDoc executable.
Test IDs (BT-*) match tests/bundle/TEST_PLAN.md. `-m smoke` selects the CI subset.
"""
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import time
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest
from conftest import (
    FEATURES,
    IS_WINDOWS,
    REPO_ROOT,
    copy_fixture,
    free_port,
    have_chrome,
    pdf_pages,
    pdf_text,
)

VERSION_RE = re.compile(r"^\d+\.\d+\.\d+", re.M)
THREE_REQS = "html2pdf/02_three_top_level_requirements"
needs_chrome = pytest.mark.skipif(not have_chrome(), reason="no Chrome/Chromium installed")


def count_nodes(sdoc_text, tag="REQUIREMENT"):
    return len(re.findall(rf"^\[{tag}\]\s*$", sdoc_text, re.M))


def field_values(sdoc_text, field):
    return re.findall(rf"^{field}: (.+)$", sdoc_text, re.M)


# ----------------------------------------------------------------------------
# BT-CLI: version, help, about, new
# ----------------------------------------------------------------------------
@pytest.mark.smoke
@pytest.mark.parametrize("args", [["version"], ["--version"], ["-v"]])
def test_cli_01_version(sd, args):
    _, out = sd.run(*args)
    assert VERSION_RE.search(out), out


def top_level_commands(sd):
    _, out = sd.run("--help")
    m = re.search(r"\{([a-z0-9,_-]+)\}", out)
    assert m, out
    return m.group(1).split(",")


@pytest.mark.smoke
def test_cli_02_help_top_level(sd):
    cmds = top_level_commands(sd)
    for expected in ["export", "server", "convert", "manage", "version"]:
        assert expected in cmds, cmds


def test_cli_02_help_every_subcommand(sd):
    cmds = top_level_commands(sd)
    failures = []
    for c in cmds:
        rc, out = sd.run(c, "--help", check=False)
        if rc != 0 or "usage:" not in out or "Traceback" in out:
            failures.append((c, rc, out[-500:]))
    _, out = sd.run("manage", "--help")
    subs = re.search(r"\{([a-z0-9,_-]+)\}", out).group(1).split(",")
    for s in subs:
        rc, o = sd.run("manage", s, "--help", check=False)
        if rc != 0 or "usage:" not in o:
            failures.append(("manage " + s, rc, o[-500:]))
    assert not failures, failures


def test_cli_03_about(sd):
    rc, out = sd.run("about")
    assert "StrictDoc" in out


def test_cli_04_new_project_then_export(sd, work):
    sd.run("new", work / "proj")
    assert list((work / "proj").rglob("*.sdoc")), "new created no .sdoc"
    sd.run("export", work / "proj", "--output-dir", work / "out", cwd=work)
    assert (work / "out" / "html" / "index.html").is_file()


def test_cli_05_bad_args_clean_error(sd, work):
    rc, out = sd.run("export", "--formats", "nope", work, check=False)
    assert rc != 0 and "Traceback" not in out, out


# ----------------------------------------------------------------------------
# BT-EXP: export to every format
# ----------------------------------------------------------------------------
def _assert_static_refs_exist(html_root, page):
    text = page.read_text(encoding="utf-8")
    refs = set(re.findall(r'(?:href|src)="([^"#?]+\.(?:css|js|svg|ico|png|woff2?|ttf))"', text))
    missing = []
    for r in refs:
        if r.startswith(("http:", "https:", "data:")):
            continue
        if not (page.parent / r).resolve().is_file():
            missing.append(r)
    assert not missing, f"{page}: missing assets {missing}"


@pytest.mark.smoke
def test_exp_01_html(sd, work):
    proj = copy_fixture(THREE_REQS, work / "project")
    sd.run("export", proj, "--formats", "html", "--output-dir", work / "out", cwd=work)
    html = work / "out" / "html"
    pages = [html / "index.html", html / "project" / "input.html",
             html / "project" / "nested" / "input2.html"]
    for p in pages:
        assert p.is_file(), p
        _assert_static_refs_exist(html, p)
    assert "Dummy high-level requirement #1" in pages[1].read_text(encoding="utf-8")


def test_exp_01b_html_strictdoc_docs_no_parallel(sd, work):
    # docs/ links into spec/ (SDOC_MARKDOWN_SPEC), so both are inputs.
    sd.run("export", REPO_ROOT / "docs", REPO_ROOT / "spec", "--no-parallelization",
           "--output-dir", work / "out", cwd=work, timeout=900)
    docs = list((work / "out" / "html").rglob("strictdoc_01_user_guide.html"))
    assert docs, "user guide not exported"
    _assert_static_refs_exist(work / "out" / "html", docs[0])


FORMAT_CASES = {
    # format: (fixture, glob of expected output, validator)
    "markdown": (THREE_REQS, "markdown/**/*.md", "text:Dummy high-level requirement #1"),
    "rst": (THREE_REQS, "rst/**/*.rst", "text:Dummy high-level requirement #1"),
    "json": (THREE_REQS, "json/**/*.json", "json"),
    "excel": (THREE_REQS, "excel/**/*.xlsx", "xlsx"),
    "reqif-sdoc": (THREE_REQS, "reqif/**/*.reqif", "xml"),
    "reqifz-sdoc": (THREE_REQS, "**/*.reqifz", "zip"),
    "sdoc": (THREE_REQS, "sdoc/**/*.sdoc", "text:[REQUIREMENT]"),
    "doxygen": ("doxygen/01_basic_export", "doxygen/**/*.tag", "xml"),
    "spdx": ("spdx/01_minimal_document_mid_enabled", "spdx/**/*.spdx*", "text:Dummy high-level requirement #1"),
}


@pytest.mark.parametrize("fmt", list(FORMAT_CASES))
def test_exp_02_formats(sd, work, fmt):
    fixture, pattern, validator = FORMAT_CASES[fmt]
    proj = copy_fixture(fixture, work / "project")
    # spdx/doxygen fixtures expect their config one level up (see their itest)
    cfg = proj / "strictdoc_config.py"
    if cfg.exists():
        shutil.copy(cfg, work / "strictdoc_config.py")
    for extra in ["file.py"]:
        if (proj / extra).exists():
            shutil.copy(proj / extra, work / extra)
    out = work / "out"
    sd.run("export", proj, "--formats", fmt, "--output-dir", out, cwd=work)
    files = [p for p in out.glob(pattern) if p.is_file()]
    if not files:  # tolerate layout differences, report what was produced
        produced = sorted(str(p.relative_to(out)) for p in out.rglob("*") if p.is_file())
        pytest.fail(f"{fmt}: no files matching {pattern}; produced: {produced[:30]}")
    for f in files:
        assert f.stat().st_size > 0, f
        if validator == "json":
            json.loads(f.read_text(encoding="utf-8"))
        elif validator == "xml":
            ET.parse(f)
        elif validator == "zip":
            assert zipfile.ZipFile(f).namelist()
        elif validator == "xlsx":
            import openpyxl
            ws = openpyxl.load_workbook(f).active
            assert ws.max_row >= 2, f"{f}: no data rows"
    if validator.startswith("text:"):
        needle = validator[5:]
        assert any(needle in f.read_text(encoding="utf-8") for f in files), needle


def test_exp_03_strictdoc_own_docs_with_config(sd, work):
    """
    StrictDoc's own repo: strictdoc_config.py importing strictdoc.api, a
    user-module statistics generator (docs.sdoc_project_statistics), source
    traceability, tree map, traceability matrix, project statistics.
    """
    proj = work / "sdroot"
    shutil.copytree(REPO_ROOT, proj, ignore=shutil.ignore_patterns(
        ".git", "build", "output", "Output", "node_modules", "__pycache__"))
    sd.run("export", ".", "--formats", "html", "--output-dir", work / "out",
           cwd=proj, timeout=1200)
    html = work / "out" / "html"
    for f in ["index.html", "project_statistics.html", "tree_map.html",
              "traceability_matrix.html", "source_coverage.html"]:
        assert (html / f).is_file(), f


def test_exp_04_diff_dirs(sd, work):
    fx = FEATURES / "diff" / "01__requirements__from_empty_doc_to_one_requirement"
    shutil.copy(fx / "strictdoc_config.py", work)
    sd.run("export", ".", "--generate-diff-dirs", fx / "lhs", fx / "rhs",
           "--output-dir", work, cwd=work)
    ch = (work / "html" / "changelog.html").read_text(encoding="utf-8")
    assert "Requirement #1" in ch and "1 (1 added)" in ch


def test_exp_05_diff_git(sd, work):
    if not shutil.which("git"):
        pytest.skip("git not installed")
    fx = FEATURES / "diff" / "01__requirements__from_empty_doc_to_one_requirement"
    repo = work / "g"
    repo.mkdir()
    g = lambda *a: subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t", *a],
                                  cwd=repo, check=True, capture_output=True)
    g("init", "-q")
    shutil.copy(fx / "strictdoc_config.py", repo)
    shutil.copytree(fx / "lhs", repo / "docs")
    g("add", "-A"); g("commit", "-qm", "lhs")
    shutil.rmtree(repo / "docs"); shutil.copytree(fx / "rhs", repo / "docs")
    g("add", "-A"); g("commit", "-qm", "rhs")
    sd.run("export", ".", "--generate-diff-git", "HEAD^..HEAD",
           "--output-dir", work / "out", cwd=repo)
    assert "Requirement #1" in (work / "out" / "html" / "changelog.html").read_text(encoding="utf-8")


# ----------------------------------------------------------------------------
# BT-CONV: import (convert) Excel and ReqIF
# ----------------------------------------------------------------------------
EXCEL_FIXTURES = ["01_import_basic_e2e", "02_float_fields_converted_to_strings",
                  "03_default_fields_are_not_duplicated"]


@pytest.mark.parametrize("name", EXCEL_FIXTURES)
@pytest.mark.parametrize("ext", ["xlsx", "xls"])
def test_conv_01_excel_import_matches_expected(sd, work, name, ext):
    fx = FEATURES / "excel" / "import" / name
    src = fx / f"input.{ext}"
    if not src.exists():
        pytest.skip(f"no {src.name} in fixture")
    expected = fx / ("expected_xslx" if ext == "xlsx" else "expected") / "expected.sdoc"
    sd.run("convert", src, work / "out", "--input-format=excel")
    got = (work / "out" / "input.sdoc").read_text(encoding="utf-8")
    assert got.replace("\r\n", "\n") == expected.read_text(encoding="utf-8").replace("\r\n", "\n")


def _excel_import_and_check(sd, work, name):
    import openpyxl
    fx = FEATURES / "excel" / "import" / name
    sd.run("convert", fx / "input.xlsx", work / "sdoc", "--input-format=excel")
    sdoc = (work / "sdoc" / "input.sdoc").read_text(encoding="utf-8")
    expected = (fx / "expected_xslx" / "expected.sdoc").read_text(encoding="utf-8")
    assert sdoc.replace("\r\n", "\n") == expected.replace("\r\n", "\n")
    ws = openpyxl.load_workbook(fx / "input.xlsx").active
    data_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
    assert count_nodes(sdoc) == len(data_rows), (count_nodes(sdoc), len(data_rows))
    header = [str(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    uid_col = [i for i, h in enumerate(header)
               if h and h.strip().upper() in ("UID", "REQ ID", "ID", "REQUIREMENT ID")]
    if uid_col:
        xl_uids = {str(r[uid_col[0]]).strip() for r in data_rows if r[uid_col[0]]}
        assert xl_uids == set(field_values(sdoc, "UID")), (xl_uids, field_values(sdoc, "UID"))
    assert "[GRAMMAR]" in sdoc
    return sdoc


def _reexport_html(sd, work, sdoc):
    sd.run("export", work / "sdoc", "--output-dir", work / "html")
    page = next((work / "html" / "html").rglob("input.html"))
    text = page.read_text(encoding="utf-8")
    for uid in field_values(sdoc, "UID"):
        assert uid in text, uid


@pytest.mark.smoke
def test_conv_02_excel_import_content(sd, work):
    """Critical path: real .xlsx (with a Parent relation column) -> .sdoc content."""
    _excel_import_and_check(sd, work, "01_import_basic_e2e")


@pytest.mark.smoke
def test_conv_02b_excel_import_then_reexport_html(sd, work):
    sdoc = _excel_import_and_check(sd, work, "03_default_fields_are_not_duplicated")
    _reexport_html(sd, work, sdoc)


# U1 (fixed on this branch): the Excel importer now declares the Parent
# relation in the generated grammar, so the imported document exports.
@pytest.mark.smoke
def test_conv_02c_excel_import_with_relations_then_reexport(sd, work):
    sdoc = _excel_import_and_check(sd, work, "01_import_basic_e2e")
    _reexport_html(sd, work, sdoc)


def test_conv_03_excel_round_trip(sd, work):
    """Export --formats excel --fields uid,title,statement, convert back, compare."""
    src = FEATURES / "excel" / "export" / "01_basic_excel_export" / "input.sdoc"
    orig = src.read_text(encoding="utf-8")
    (work / "in").mkdir()
    shutil.copy(src, work / "in" / "input.sdoc")
    sd.run("export", work / "in", "--formats", "excel", "--fields", "uid,title,statement",
           "--output-dir", work / "x")
    xlsx = next((work / "x").rglob("*.xlsx"))
    sd.run("convert", xlsx, work / "back", "--input-format=excel")
    back = next((work / "back").glob("*.sdoc")).read_text(encoding="utf-8")
    assert count_nodes(back) == count_nodes(orig), (count_nodes(back), count_nodes(orig))
    assert set(field_values(back, "UID")) == set(field_values(orig, "UID"))
    for t in field_values(orig, "TITLE")[1:]:  # [0] is the document title
        assert t in back, t
    for st in field_values(orig, "STATEMENT"):
        assert st in back, st
    sd.run("export", work / "back", "--output-dir", work / "html2")
    assert list((work / "html2" / "html").rglob("*.html"))


# U1 (fixed on this branch): the default Excel round trip works.
def test_conv_03b_excel_round_trip_default_fields(sd, work):
    src = FEATURES / "excel" / "export" / "01_basic_excel_export" / "input.sdoc"
    (work / "in").mkdir()
    shutil.copy(src, work / "in" / "input.sdoc")
    sd.run("export", work / "in", "--formats", "excel", "--output-dir", work / "x")
    sd.run("convert", next((work / "x").rglob("*.xlsx")), work / "back", "--input-format=excel")
    sd.run("export", work / "back", "--output-dir", work / "html2")


def test_conv_04_reqif_import(sd, work):
    fx = FEATURES / "reqif" / "profiles" / "p01_sdoc" / "examples" / "01_sample" / "sample.reqif"
    sd.run("convert", fx, work / "out")
    sdoc = (work / "out" / "sample.sdoc").read_text(encoding="utf-8")
    assert "[REQUIREMENT_TYPE]" in sdoc or count_nodes(sdoc) > 0
    sd.run("export", work / "out", "--output-dir", work / "html")
    assert list((work / "html" / "html").rglob("sample.html"))


def test_conv_05_reqif_round_trip(sd, work):
    proj = copy_fixture(THREE_REQS, work / "project")
    sd.run("export", proj, "--formats", "reqif-sdoc", "--output-dir", work / "r", cwd=work)
    reqif = next((work / "r").rglob("*.reqif"))
    sd.run("convert", reqif, work / "back")
    back = "\n".join(p.read_text(encoding="utf-8") for p in (work / "back").rglob("*.sdoc"))
    for t in ["Dummy high-level requirement #1", "Dummy high-level requirement #3"]:
        assert t in back, t


# ----------------------------------------------------------------------------
# BT-MAN / BT-FMT: manage auto-uid, assets, new; format
# ----------------------------------------------------------------------------
def test_man_01_auto_uid(sd, work):
    d = work / "p"; d.mkdir()
    (d / "doc.sdoc").write_text(
        "[DOCUMENT]\nTITLE: T\n\n[REQUIREMENT]\nTITLE: A\nSTATEMENT: a\n\n"
        "[REQUIREMENT]\nTITLE: B\nSTATEMENT: b\n", encoding="utf-8")
    sd.run("manage", "auto-uid", d, cwd=work)
    uids = field_values((d / "doc.sdoc").read_text(encoding="utf-8"), "UID")
    assert len(uids) == 2 and len(set(uids)) == 2, uids


def test_man_02_assets(sd, work):
    proj = copy_fixture(THREE_REQS, work / "p")
    sd.run("manage", "assets", proj, cwd=work)


def test_man_03_new_node(sd, work):
    d = work / "p"; d.mkdir()
    (d / "doc.sdoc").write_text("[DOCUMENT]\nTITLE: T\n\n[REQUIREMENT]\nUID: REQ-1\nSTATEMENT: a\n",
                                encoding="utf-8")
    sd.run("manage", "new", d, cwd=work)
    assert count_nodes((d / "doc.sdoc").read_text(encoding="utf-8")) == 2


def test_fmt_01_format(sd, work):
    proj = copy_fixture(THREE_REQS, work / "p")
    sd.run("format", proj, cwd=work)
    assert "Dummy high-level requirement #1" in (proj / "input.sdoc").read_text(encoding="utf-8")


# ----------------------------------------------------------------------------
# BT-H2P: html2pdf (critical)
# ----------------------------------------------------------------------------
@pytest.mark.smoke
@needs_chrome
def test_h2p_01_export_pdf_content(sd, work, results_dir):
    proj = copy_fixture(THREE_REQS, work / "project")
    rc, out = sd.run("export", proj, "--formats", "html2pdf", "--output-dir",
                     work / "out", cwd=work, timeout=900)
    (results_dir / "h2p_01.log").write_text(out, encoding="utf-8")
    pdfs = {p.relative_to(work / "out" / "html2pdf" / "pdf").as_posix(): p
            for p in (work / "out" / "html2pdf" / "pdf").rglob("*.pdf")}
    assert set(pdfs) >= {"input.pdf", "nested/input2.pdf", "nested/subnested/input3.pdf"}, pdfs
    for name, p in pdfs.items():
        assert p.read_bytes()[:4] == b"%PDF", name
        assert pdf_pages(p) >= 1, name
        text = pdf_text(p)
        assert "Dummy" in text, f"{name}: no requirement text extracted"
    t = pdf_text(pdfs["input.pdf"])
    for needle in ["Dummy Software Requirements", "Dummy high-level requirement #1",
                   "Dummy high-level requirement #3"]:
        assert needle in " ".join(t.split()), needle
    # how the driver was located (recorded for the report)
    m = re.findall(r"html2pdf4doc: (Installed Chrome version.*|ChromeDriver (?:exists|available|downloaded).*)", out)
    (results_dir / "h2p_driver_discovery.txt").write_text("\n".join(m), encoding="utf-8")


@needs_chrome
def test_h2p_02_bad_chrome_binary_clean_error(sd, work):
    proj = copy_fixture(THREE_REQS, work / "project")
    rc, out = sd.run("export", proj, "--formats", "html2pdf", "--chrome-binary",
                     work / "no-such-chrome", "--output-dir", work / "out",
                     cwd=work, check=False)
    assert rc != 0
    assert "chrome" in out.lower() and "Traceback" not in out, out[-2000:]


@needs_chrome
def test_h2p_03_offline_without_cached_driver(sd, work):
    """No network + empty driver cache -> must fail with a clear message, not hang."""
    if not shutil.which("unshare") or IS_WINDOWS:
        pytest.skip("needs `unshare -rn` (Linux user namespaces)")
    probe = subprocess.run(["unshare", "-rn", "true"], capture_output=True)
    if probe.returncode != 0:
        pytest.skip("user namespaces unavailable")
    proj = copy_fixture(THREE_REQS, work / "project")
    cfg = proj / "strictdoc_config.py"
    cfg.write_text(cfg.read_text().replace("./Output/cache", str(work / "emptycache")))
    from conftest import clean_env
    t0 = time.time()
    proc = subprocess.run(offline_cmd([*sd.cmd, "export", str(proj), "--formats",
                           "html2pdf", "--output-dir", str(work / "out")]),
                          cwd=work, capture_output=True, text=True, timeout=600,
                          env=clean_env())
    out = proc.stdout + proc.stderr
    assert proc.returncode != 0, "html2pdf unexpectedly succeeded offline"
    assert time.time() - t0 < 300, "offline failure took too long"
    # Must tell the user *why*: the ChromeDriver download failed. (A raw
    # html2pdf4doc traceback is printed today, same as with pip strictdoc:
    # tracked as UX issue U3, not failed here.)
    assert re.search(r"GET request failed|ChromeDriver|chromedriver", out), out[-3000:]


@needs_chrome
def test_h2p_04_offline_with_cached_driver(sd, work):
    """Driver already cached: html2pdf must work with no network."""
    if not shutil.which("unshare") or IS_WINDOWS:
        pytest.skip("needs `unshare -rn`")
    if subprocess.run(["unshare", "-rn", "true"], capture_output=True).returncode != 0:
        pytest.skip("user namespaces unavailable")
    proj = copy_fixture(THREE_REQS, work / "project")
    cfg = proj / "strictdoc_config.py"
    cache = work / "cache"
    cfg.write_text(cfg.read_text().replace("./Output/cache", str(cache)))
    sd.run("export", proj, "--formats", "html2pdf", "--output-dir", work / "o1", cwd=work)
    from conftest import clean_env
    proc = subprocess.run(offline_cmd([*sd.cmd, "export", str(proj), "--formats",
                           "html2pdf", "--output-dir", str(work / "o2")]),
                          cwd=work, capture_output=True, text=True, timeout=600, env=clean_env())
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]
    assert (work / "o2" / "html2pdf" / "pdf" / "input.pdf").is_file()


# ----------------------------------------------------------------------------
# BT-SRV: server (boot, headless crawl, server-side actions)
# ----------------------------------------------------------------------------
FULL_FEATURE_CONFIG = """from strictdoc.core.project_config import ProjectConfig


def create_config() -> ProjectConfig:
    return ProjectConfig(
        dir_for_sdoc_cache="./Output/cache",
        project_features=[
            "TABLE_SCREEN", "TRACEABILITY_SCREEN", "DEEP_TRACEABILITY_SCREEN",
            "SEARCH", "HTML2PDF", "PROJECT_STATISTICS_SCREEN", "TREE_MAP_SCREEN",
            "TRACEABILITY_MATRIX_SCREEN", "DIFF", "MATHJAX", "MERMAID", "NESTOR",
        ],
    )
"""


def write_full_feature_config(proj):
    (proj / "strictdoc_config.py").write_text(FULL_FEATURE_CONFIG)


def offline_cmd(cmd):
    """Run cmd with no network but a working loopback (chromedriver needs it)."""
    import shlex
    inner = "ip link set lo up 2>/dev/null; exec " + " ".join(shlex.quote(c) for c in cmd)
    return ["unshare", "-rn", "sh", "-c", inner]


class Server:
    def __init__(self, sd, project, work, extra=()):
        self.port = free_port()
        self.base = f"http://127.0.0.1:{self.port}"
        self.log = work / f"server-{self.port}.log"
        self.proc = sd.popen("server", project, "--port", self.port, *extra,
                             cwd=project, log_path=self.log)

    def wait(self, timeout=180):
        t0 = time.time()
        while time.time() - t0 < timeout:
            if self.proc.poll() is not None:
                raise AssertionError("server exited:\n" + self.log.read_text()[-4000:])
            try:
                urllib.request.urlopen(self.base + "/", timeout=5)
                return self
            except Exception:
                time.sleep(1)
        raise AssertionError("server did not come up:\n" + self.log.read_text()[-4000:])

    def stop(self):
        if self.proc.poll() is None:
            if IS_WINDOWS:
                self.proc.terminate()
            else:
                self.proc.send_signal(signal.SIGINT)
            try:
                self.proc.wait(20)
            except subprocess.TimeoutExpired:
                self.proc.kill()
        log = self.log.read_text(errors="replace")
        return log


def _crawl(server, results_dir, name, *args):
    out = results_dir / f"crawl_{name}.json"
    proc = subprocess.run([sys.executable, str(Path(__file__).parent / "server_crawl.py"),
                           "--base-url", server.base, "--out", str(out), *args],
                          capture_output=True, text=True, timeout=3600)
    (results_dir / f"crawl_{name}.log").write_text(proc.stdout + proc.stderr)
    return proc.returncode, json.loads(out.read_text()) if out.exists() else None, proc.stdout + proc.stderr


@pytest.mark.smoke
def test_srv_01_boot_and_smoke_crawl(sd, work, results_dir):
    pytest.importorskip("playwright")
    proj = copy_fixture(THREE_REQS, work / "project")
    write_full_feature_config(proj)
    s = Server(sd, proj, work).wait()
    try:
        rc, data, out = _crawl(s, results_dir, "smoke", "--smoke")
    finally:
        log = s.stop()
    assert "Traceback" not in log, log[-4000:]
    assert data is not None, out
    bad = [e for e in data["pages"] + data["searches"] + data["forms"] if e["fail"]]
    assert not bad, json.dumps(bad, indent=1)[:6000]


def test_srv_02_full_crawl_strictdoc_docs(sd, work, results_dir):
    """Full crawl + open edit/create forms on a scratch copy of StrictDoc's repo."""
    pytest.importorskip("playwright")
    proj = work / "sdroot"
    shutil.copytree(REPO_ROOT, proj, ignore=shutil.ignore_patterns(
        ".git", "build", "output", "Output", "node_modules", "__pycache__"))
    s = Server(sd, proj, work).wait(600)
    try:
        ignore = os.environ.get("BUNDLE_TEST_CRAWL_IGNORE_CONSOLE")
        extra = ["--forms", "--max-pages", os.environ.get("BUNDLE_TEST_MAX_PAGES", "250"),
                 # UPSTREAM U4 (same with pip strictdoc): free-text search results
                 # render document images with document-relative "_assets/x.png"
                 # URLs, which 404 when resolved against /search.
                 "--known-upstream", r"/search\?\S* \S+/_assets/"]
        if ignore:
            extra += ["--ignore-console", ignore]
        rc, data, out = _crawl(s, results_dir, "full", *extra)
    finally:
        log = s.stop()
    (results_dir / "server_full.log").write_text(log)
    assert data is not None, out
    bad = [e for e in data["pages"] + data["searches"] + data["forms"] if e["fail"]]
    assert not bad, json.dumps(bad, indent=1)[:8000]
    assert "Traceback" not in log, log[-4000:]


@needs_chrome
def test_srv_03_server_html2pdf_endpoint(sd, work, results_dir):
    """The server's 'Export to PDF' button: html2pdf from inside the server process."""
    proj = copy_fixture(THREE_REQS, work / "project")
    s = Server(sd, proj, work).wait()
    try:
        html = urllib.request.urlopen(s.base + "/project/input.html", timeout=60).read().decode()
        m = re.search(r'href="(/export_html2pdf/[0-9a-f]+)"', html)
        assert m, "no export_html2pdf link on document page"
        resp = urllib.request.urlopen(s.base + m.group(1), timeout=600)
        body = resp.read()
    finally:
        log = s.stop()
    (results_dir / "server_h2p.log").write_text(log)
    assert body[:4] == b"%PDF", body[:300]


def test_srv_04_server_reqif_export_endpoint(sd, work):
    proj = copy_fixture(THREE_REQS, work / "project")
    s = Server(sd, proj, work).wait()
    try:
        body = urllib.request.urlopen(s.base + "/reqif/export_tree", timeout=120).read()
    finally:
        s.stop()
    assert b"REQ-IF" in body[:4000] or body[:2] == b"PK", body[:300]


def test_srv_05_server_watch_mode_boots(sd, work):
    proj = copy_fixture(THREE_REQS, work / "project")
    s = Server(sd, proj, work, extra=("--watch",)).wait()
    try:
        page = proj / "input.sdoc"
        page.write_text(page.read_text().replace("Dummy high-level requirement #1",
                                                 "Edited requirement #1"))
        time.sleep(4)
        html = urllib.request.urlopen(s.base + "/project/input.html", timeout=60).read().decode()
    finally:
        log = s.stop()
    assert "Traceback" not in log, log[-4000:]
    assert "Edited requirement #1" in html


# ----------------------------------------------------------------------------
# BT-BND: bundle-specific behaviour
# ----------------------------------------------------------------------------
def test_bnd_02_python_env_is_ignored(sd, work):
    if not sd.is_bundle:
        pytest.skip("bundle only")
    proj = copy_fixture(THREE_REQS, work / "project")
    sd.run("export", proj, "--output-dir", work / "out", cwd=work,
           env_extra={"PYTHONHOME": str(work / "bogus"), "PYTHONPATH": str(work / "bogus"),
                      "PYTHONSAFEPATH": "1"})
    assert (work / "out" / "html" / "index.html").is_file()


def test_bnd_03_relocated_bundle_path_with_spaces(sd, work):
    if not sd.is_bundle:
        pytest.skip("bundle only")
    dst = work / "dir with spaces" / "strictdoc bundle"
    shutil.copytree(sd.bundle_path.parent, dst, symlinks=True)
    exe = dst / sd.bundle_path.name
    proj = copy_fixture(THREE_REQS, work / "project")
    from conftest import clean_env
    proc = subprocess.run([str(exe), "export", str(proj), "--output-dir", str(work / "out")],
                          cwd=work, capture_output=True, text=True, env=clean_env(), timeout=600)
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]
    assert (work / "out" / "html" / "index.html").is_file()


def test_bnd_04_launcher_starts(sd, work):
    """Experimental Tk launcher: frozen tkinter must import; window must open."""
    if not sd.is_bundle:
        pytest.skip("bundle only")
    if not IS_WINDOWS and not os.environ.get("DISPLAY"):
        pytest.skip("no DISPLAY")
    log = work / "launcher.log"
    p = sd.popen("launcher", cwd=work, log_path=log,
                 env_extra={"HOME": str(work), "DISPLAY": os.environ.get("DISPLAY", "")})
    time.sleep(10)
    alive = p.poll() is None
    if alive:
        p.kill(); p.wait()
    text = log.read_text(errors="replace")
    assert alive, f"launcher exited early ({p.returncode}):\n{text[-3000:]}"
    assert "Traceback" not in text, text[-3000:]
